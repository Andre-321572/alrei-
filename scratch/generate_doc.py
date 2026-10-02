import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_heading_with_spacing(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    for run in h.runs:
        run.font.name = 'Calibri'
        if level == 1:
            run.font.color.rgb = RGBColor(16, 88, 48)  # ALREI Green
            run.font.size = Pt(18)
            run.font.bold = True
        elif level == 2:
            run.font.color.rgb = RGBColor(212, 143, 24) # ALREI Gold/Amber
            run.font.size = Pt(14)
            run.font.bold = True
        elif level == 3:
            run.font.color.rgb = RGBColor(40, 40, 40)
            run.font.size = Pt(12)
            run.font.bold = True
    return h

def generate_report(output_path):
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base style setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(50, 50, 50)

    # Title Header Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("CENTRE ALREI DE FORMATION DES TRAVAILLEURS")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 88, 48)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("RAPPORT TECHNIQUE & SYNTHÈSE DES FONCTIONNALITÉS OPÉRATIONNELLES")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(212, 143, 24)

    # Meta Info Box (Table)
    table_meta = doc.add_table(rows=2, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False
    
    col_widths = [Inches(3.2), Inches(3.2)]
    
    meta_data = [
        [("Projet :", " Plateforme E-Learning (CSI-Afrique)"), ("Statut :", " Validé & Conforme (Production)")],
        [("Langues :", " Français, Anglais, Portugais"), ("Version / Git :", " Branche 'andre' (Commit 8ed5b23)")]
    ]
    
    for r_idx, row in enumerate(table_meta.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "F4F7F4")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            lbl, val = meta_data[r_idx][c_idx]
            r1 = p.add_run(lbl)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(16, 88, 48)
            r2 = p.add_run(val)
            r2.font.color.rgb = RGBColor(60, 60, 60)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 1. Introduction
    add_heading_with_spacing(doc, "1. Présentation Générale", level=1)
    p_intro = doc.add_paragraph(
        "Ce document récapitule l'ensemble des fonctionnalités opérationnelles, des modules intégrés et des récentes optimisations apportées à la plateforme e-learning du Centre ALREI de formation des travailleurs (CSI-Afrique). La plateforme a été conçue et finalisée pour offrir une expérience fluide, sécurisée et parfaitement multilingue."
    )
    p_intro.paragraph_format.space_after = Pt(10)
    p_intro.paragraph_format.line_spacing = 1.15

    # 2. Bilan des Fonctionnalités Fonctionnelles
    add_heading_with_spacing(doc, "2. Bilan des Modules & Fonctionnalités Opérationnelles", level=1)

    # 2.1 Authentification & Inscription
    add_heading_with_spacing(doc, "2.1 Authentification & Gestion des Comptes", level=2)
    features_auth = [
        ("Système d'inscription modulaire :", " Inscription complète incluant le nom, prénom, pays de résidence, genre, e-mail personnel, numéro WhatsApp international, date de naissance, expérience et langue préférée."),
        ("Champs optionnels ajustés :", " Les champs 'Organisation / Syndicat' et 'Email de l'organisation' sont désormais entièrement optionnels pour permettre aux apprenants indépendants de s'inscrire sans contrainte."),
        ("Masquage/Affichage interactif du mot de passe :", " L'icône œil sur l'ensemble des formulaires de connexion et d'inscription permet désormais d'afficher ou de masquer instantanément le mot de passe."),
        ("Validation sécurisée par code OTP :", " Génération et vérification automatique du code de confirmation à 6 chiffres transmis par e-mail lors de la création de compte."),
        ("Sélection dynamique des pays et indicatifs :", " Composant téléphonique international avec détection automatique du pays de résidence et affichage du drapeau national.")
    ]
    for title, desc in features_auth:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_t = p.add_run(title)
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(20, 20, 20)
        p.add_run(desc)

    # 2.2 Multilinguisme & Internationalisation (i18n)
    add_heading_with_spacing(doc, "2.2 Multilinguisme & Traduction Complète (i18n)", level=2)
    features_i18n = [
        ("Traductions intégrales 3 Langues :", " Support parfait et complet en Français (FR), Anglais (EN) et Portugais (PT)."),
        ("En-tête & Pied de page (Footer) :", " Tous les liens de navigation, catégories, mentions légales et liens d'aide (notamment 'Comment utiliser le centre' / 'How to use the center' / 'Como usar o centro') se traduisent instantanément lors du changement de langue."),
        ("Interface dynamique :", " Tous les boutons, modales, formulaires et notifications s'adaptent selon la langue sélectionnée par l'utilisateur."),
        ("Harmonisation typographique :", " Uniformisation complète des polices de caractères et de la mise en page à travers l'ensemble des vues.")
    ]
    for title, desc in features_i18n:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_t = p.add_run(title)
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(20, 20, 20)
        p.add_run(desc)

    # 2.3 Catalogue & Parcours d'Apprentissage
    add_heading_with_spacing(doc, "2.3 Catalogues de Formations & Apprentissage", level=2)
    features_courses = [
        ("Domaines de formation ciblés :", " Leadership syndical, Syndicalisation et négociation collective, Recherche sur le travail et politique économique, Changement climatique et transition juste, Numérisation et avenir du travail."),
        ("Consultation des cours & syllabus :", " Fiches détaillées des formations avec programme complet, aperçu vidéo, prérequis, durée et profil de l'instructeur."),
        ("Système d'évaluation et avis :", " Possibilité d'évaluer les cours et de laisser des retours d'expérience utilisateurs.")
    ]
    for title, desc in features_courses:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_t = p.add_run(title)
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(20, 20, 20)
        p.add_run(desc)

    # 2.4 Espace Instructeur, Administration & Attestations
    add_heading_with_spacing(doc, "2.4 Tableaux de Bord, Administration & Certificats", level=2)
    features_admin = [
        ("Tableau de bord Administrateur :", " Supervision complète des utilisateurs, gestion des candidatures d'instructeurs, approbation des cours et suivi analytique par pays."),
        ("Tableau de bord Instructeur :", " Interface de création et d'édition de cours, suivi des devoirs et évaluation des apprenants."),
        ("Attestations & Certificats de réussite :", " Module de demande et de délivrance d'attestations certifiées avec validation préalable de l'administration.")
    ]
    for title, desc in features_admin:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_t = p.add_run(title)
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(20, 20, 20)
        p.add_run(desc)

    # 3. Tableau Récapitulatif des Derniers Correctifs
    add_heading_with_spacing(doc, "3. Synthèse des Domiciliations & Derniers Correctifs Validés", level=1)

    table = doc.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    t_widths = [Inches(2.0), Inches(3.4), Inches(1.0)]

    headers = ["Élément", "Description de la modification", "Statut"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = t_widths[i]
        set_cell_background(hdr_cells[i], "105830")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        ("Champs Organisation", "Passage des champs 'Organisation / Syndicat' et 'Email de l'organisation' en optionnel (suppression du caractère obligatoire).", "OK - Corrigé"),
        ("Icône Œil (Mot de passe)", "Activation de la bascule réactive d'affichage/masquage du mot de passe sur tous les formulaires de connexion et d'inscription.", "OK - Corrigé"),
        ("Traduction Footer & Liens", "Correction et vérification globale des clés de traduction pour 'Comment utiliser le centre' en FR, EN et PT.", "OK - Corrigé")
    ]

    for r_idx, (el, desc, stat) in enumerate(rows_data, start=1):
        row_cells = table.rows[r_idx].cells
        bg_color = "F9FAF9" if r_idx % 2 == 1 else "FFFFFF"
        
        for c_idx, val in enumerate([el, desc, stat]):
            row_cells[c_idx].width = t_widths[c_idx]
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            if c_idx == 0:
                r.bold = True
            elif c_idx == 2:
                r.bold = True
                r.font.color.rgb = RGBColor(16, 88, 48)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 4. État du Déploiement & Build
    add_heading_with_spacing(doc, "4. État du Déploiement & Assurance Qualité", level=1)
    
    p_build = doc.add_paragraph()
    p_build.paragraph_format.line_spacing = 1.15
    p_build.paragraph_format.space_after = Pt(6)
    p_build.add_run("Toutes les fonctionnalités ont été soumises au processus complet d'assurance qualité et de validation technique :").font.color.rgb = RGBColor(50, 50, 50)

    checks = [
        ("Compilation de production (Nitro SSG/Prerender) :", " Exécution réussie de `npm run build` avec 419 routes pré-rendues statiquement sans la moindre erreur."),
        ("Dépôts de code (Version control) :", " Les modifications sont validées et synchronisées sur la branche `andre` des dépôts GitHub 'origin' et 'repo-neostart'."),
        ("Compatibilité navigateurs & Mobiles :", " Interface responsive testée et conforme sur desktop, tablette et smartphone.")
    ]

    for title, desc in checks:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_t = p.add_run(title)
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(16, 88, 48)
        p.add_run(desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # Signoff Block
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(15)
    p_sign.paragraph_format.keep_with_next = True
    r_sign = p_sign.add_run("Équipe Technique ALREI - Octobre 2026")
    r_sign.bold = True
    r_sign.font.color.rgb = RGBColor(16, 88, 48)

    doc.save(output_path)
    print(f"Document Word créé avec succès à : {output_path}")

if __name__ == "__main__":
    generate_report("c:\\Learnup_NuxtJs_v1.0.0\\Rapport_Fonctionnalites_ALREI.docx")

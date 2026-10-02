import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_heading_with_spacing(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    for run in h.runs:
        run.font.name = 'Calibri'
        if level == 1:
            run.font.color.rgb = RGBColor(16, 88, 48)  # ALREI Green
            run.font.size = Pt(16)
            run.font.bold = True
        elif level == 2:
            run.font.color.rgb = RGBColor(212, 143, 24) # ALREI Gold
            run.font.size = Pt(13)
            run.font.bold = True
    return h

def generate_sept23_report(output_path):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(50, 50, 50)

    # Title Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("CENTRE ALREI DE FORMATION DES TRAVAILLEURS")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 88, 48)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("RAPPORT DÉTAILLÉ DES RÉALISATIONS (DU 23 SEPTEMBRE AU 01 OCTOBRE 2026)")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(212, 143, 24)

    # Meta Table
    table_meta = doc.add_table(rows=2, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False
    col_widths = [Inches(3.2), Inches(3.2)]
    
    meta_data = [
        [("Période couverte :", " 23 Septembre 2026 – 01 Octobre 2026"), ("Projet :", " E-Learning ALREI (CSI-Afrique)")],
        [("Environnement :", " Nuxt 3 / Vue 3 / Vite / Nitro"), ("Version Git :", " Branche 'andre' (Dépôts origin & repo-neostart)")]
    ]
    
    for r_idx, row in enumerate(table_meta.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "F4F7F4")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            lbl, val = meta_data[r_idx][c_idx]
            r1 = p.add_run(lbl)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(16, 88, 48)
            p.add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. Authentification & Sécurité
    add_heading_with_spacing(doc, "1. Authentification, Inscription & Sécurité", level=1)
    items_auth = [
        ("Vérification OTP par E-mail :", " Mise en place du mécanisme de confirmation d'inscription avec validation de code OTP à 6 chiffres."),
        ("Champs Optionnels :", " Assouplissement du formulaire d'inscription en rendant les champs 'Organisation / Syndicat' et 'Email de l'organisation' optionnels avec indicateur visuel."),
        ("Bascule Afficher/Masquer Mot de passe :", " Ajout d'une icône œil dynamique (bi-eye / bi-eye-slash) permettant d'afficher ou masquer le mot de passe sur la page d'inscription et la modale LoginModal."),
        ("Composant Téléphonique International :", " Intégration de l'InternationalPhoneInput avec détection automatique du pays et gestion dynamique des indicatifs et drapeaux.")
    ]
    for t, d in items_auth:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(t)
        r.bold = True
        p.add_run(d)

    # 2. Internationalisation & Traduction
    add_heading_with_spacing(doc, "2. Internationalisation & Multilinguisme (i18n)", level=1)
    items_i18n = [
        ("Localisation Complète en 3 Langues :", " Traduction intégrale en Français (FR), Anglais (EN) et Portugais (PT) de toutes les pages, composants, navbars, modales et dashboards."),
        ("Harmonisation du Footer & Liens d'Aide :", " Traduction systématique des liens d'assistance (ex: 'Comment utiliser le centre' / 'How to use the center' / 'Como usar o centro')."),
        ("Localisation des Crédits :", " Ajout des mentions d'édition multilingues 'Développé par / Developed by / Desenvolvido por'.")
    ]
    for t, d in items_i18n:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(t)
        r.bold = True
        p.add_run(d)

    # 3. Formations & Administration
    add_heading_with_spacing(doc, "3. Gestion des Cours & Espaces d'Administration", level=1)
    items_admin = [
        ("Administration des Formations & Formateurs :", " Enrichissement du panneau d'administration pour permettre l'assignation des instructeurs par l'admin et la gestion des notifications."),
        ("Bouton Ajouter un Module :", " Rétablissement de la fonctionnalité d'ajout de modules de cours dans la modale d'administration."),
        ("Création de Cours Enrichie :", " Amélioration du formulaire instructeur pour la structuration des cours (syllabus, prérequis, vidéos, documents)."),
        ("Modes d'Accès aux Formations :", " Support des formations gratuites, sur nomination syndicale et sur tarification standard / bourse.")
    ]
    for t, d in items_admin:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(t)
        r.bold = True
        p.add_run(d)

    # 4. Ergonomie, Design & Performance
    add_heading_with_spacing(doc, "4. Design, Ergonomie & Performance Technique", level=1)
    items_ui = [
        ("Police Typographique Unifiée :", " Application globale de la police de caractères Outfit sur l'ensemble de la plateforme."),
        ("Filtres & Recherche de Cours :", " Ajout d'un widget de recherche et d'onglets de filtrage (ITCILO) sur la page des cours (`courses.vue`)."),
        ("Build de Production Validé :", " Compilation réussie avec Nitro (419 routes pré-rendues statiquement sans erreur)."),
        ("Déploiement Git Synchronisé :", " Publication sur la branche `andre` vers les deux dépôts distants (origin et repo-neostart).")
    ]
    for t, d in items_ui:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(t)
        r.bold = True
        p.add_run(d)

    doc.add_paragraph().paragraph_format.space_after = Pt(15)
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(10)
    p_sign.paragraph_format.keep_with_next = True
    r_sign = p_sign.add_run("Rapport préparé par l'équipe technique - Octobre 2026")
    r_sign.bold = True
    r_sign.font.color.rgb = RGBColor(16, 88, 48)

    doc.save(output_path)
    print(f"Rapport enregistré à : {output_path}")

if __name__ == "__main__":
    generate_sept23_report("c:\\Learnup_NuxtJs_v1.0.0\\Rapport_Realisations_Depuis_23_Septembre.docx")
